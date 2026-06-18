import sys
import re
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def markdown_to_docx(md_path, docx_path):
    document = Document()
    
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()
        
    lines = content.split('\n')
    
    in_code_block = False
    
    for line in lines:
        # Handle code blocks
        if line.startswith('```'):
            in_code_block = not in_code_block
            continue
            
        if in_code_block:
            p = document.add_paragraph()
            run = p.add_run(line)
            run.font.name = 'Courier New'
            continue
            
        # Handle headers
        if line.startswith('# '):
            document.add_heading(line[2:].strip(), level=1)
        elif line.startswith('## '):
            document.add_heading(line[3:].strip(), level=2)
        elif line.startswith('### '):
            document.add_heading(line[4:].strip(), level=3)
        elif line.startswith('---'):
            document.add_page_break()
        elif line.strip() == '':
            document.add_paragraph('')
        else:
            # Handle list items
            if line.startswith('* ') or line.startswith('- '):
                p = document.add_paragraph(style='List Bullet')
                text = line[2:].strip()
            elif re.match(r'^\d+\.\s', line):
                p = document.add_paragraph(style='List Number')
                text = re.sub(r'^\d+\.\s', '', line).strip()
            else:
                p = document.add_paragraph()
                text = line
            
            # Simple bold/italic parsing for the rest of the text
            # split by **
            parts = text.split('**')
            for i, part in enumerate(parts):
                if i % 2 == 1:
                    run = p.add_run(part)
                    run.bold = True
                else:
                    # split by * for italic
                    subparts = part.split('*')
                    for j, subpart in enumerate(subparts):
                        if j % 2 == 1:
                            run = p.add_run(subpart)
                            run.italic = True
                        else:
                            # split by ` for code
                            codes = subpart.split('`')
                            for k, code in enumerate(codes):
                                if k % 2 == 1:
                                    run = p.add_run(code)
                                    run.font.name = 'Courier New'
                                else:
                                    p.add_run(code)

    # Save the document
    document.save(docx_path)
    print(f"Docx file saved at: {docx_path}")

if __name__ == "__main__":
    md_file = "/home/haquangchuong/Desktop/NCKH_chuong/remote-hardware-lab/docs/HUONG_DAN_SU_DUNG.md"
    docx_file = "/home/haquangchuong/Desktop/NCKH_chuong/remote-hardware-lab/docs/HUONG_DAN_SU_DUNG.docx"
    markdown_to_docx(md_file, docx_file)
